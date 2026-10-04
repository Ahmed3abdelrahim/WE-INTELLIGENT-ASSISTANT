#!/usr/bin/env python3
"""Generates slides/we_assistant.pptx (SPEC.md section 11, 10 slides). Uses only real data
gathered during this build session — see docs/progress.md and eval/results.md for sources.
Run `python scripts/make_slides.py` to (re)generate after eval/results.md is updated.
"""
import os

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_PATH = os.path.join(REPO_ROOT, "slides", "we_assistant.pptx")

PURPLE = RGBColor(0x6B, 0x3F, 0xA0)
DARK = RGBColor(0x24, 0x1C, 0x2E)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]


def add_slide(title, bullets, note=None):
    slide = prs.slides.add_slide(BLANK)
    # title
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.0))
    tf = tb.text_frame
    tf.text = title
    tf.paragraphs[0].font.size = Pt(32)
    tf.paragraphs[0].font.bold = True
    tf.paragraphs[0].font.color.rgb = PURPLE
    # accent bar
    bar = slide.shapes.add_shape(1, Inches(0.6), Inches(1.25), Inches(2.0), Pt(4))
    bar.fill.solid()
    bar.fill.fore_color.rgb = PURPLE
    bar.line.fill.background()
    # body
    body = slide.shapes.add_textbox(Inches(0.6), Inches(1.6), Inches(12.1), Inches(5.3))
    btf = body.text_frame
    btf.word_wrap = True
    for i, b in enumerate(bullets):
        p = btf.paragraphs[0] if i == 0 else btf.add_paragraph()
        indent = b.startswith("  ")
        p.text = b.strip()
        p.font.size = Pt(16 if not indent else 14)
        p.font.color.rgb = DARK
        p.level = 1 if indent else 0
    if note:
        slide.notes_slide.notes_text_frame.text = note
    return slide


# 1. Requirements
add_slide("Requirements", [
    "On-prem, bilingual (Arabic / English / Egyptian dialect) RAG assistant for WE Telecom Egypt",
    "Voice or text in, grounded + cited answer out; voice questions also get a spoken reply",
    "Users can upload documents (PDF/DOCX/TXT/HTML/images) to query",
    "Build constraint: 2-day case-study PoC on CPU-only laptop hardware; later verified on an NVIDIA RTX 3080",
    "Default must stay fully local; cloud LLM is opt-in, comparison-only, never the demo default",
])

# 2. Architecture
add_slide("Architecture: 6 Containers", [
    "frontend (nginx) — static UI + reverse proxy, the only container with a published host port",
    "backend (FastAPI) — guardrails, routing, RAG, ingestion, OCR, embeddings + reranker, SQLite",
    "llm (llama.cpp server) — Qwen3-4B Q4_K_M GGUF, 4 parallel slots x 4096 tokens",
    "asr (faster-whisper) — large-v3-turbo (int8 on CPU, fp16 on GPU)",
    "tts (Piper) — English + Arabic voices",
    "qdrant — hybrid dense+sparse vector search",
    "All inter-service traffic on an internal-only network",
    "  Verified as native host processes on the same ports (no Docker daemon on either test machine); compose files config-validated",
])

# 3. Ingestion and Arabic PDF handling
add_slide("Ingestion and Arabic PDF Handling", [
    "te.eg crawled: 106 pages (95 Arabic, 11 English), robots.txt-respecting, 1 req/s; isolated sessions per language (a shared cookie jar flipped Arabic pages to English)",
    "PDF: PyMuPDF per page → garbled-Arabic detection → 300 DPI render + Tesseract ara+eng OCR fallback",
    "Scanned Arabic and English PDFs, DOCX tables, TXT (cp1256 fallback), HTML, images — all answerable (heavy test)",
    "Price tables with merged header cells (colspan/rowspan) expanded into per-column headers — fixed 'col3: 775' garbage",
    "Chunking: ~450 tokens, FAQ pairs and tables kept whole when they fit, oversized ones split on rows/sentences — 360 chunks, max 509 tokens",
    "Uploads: injection sentences stripped, originals saved as data/uploads/<uuid>, private to the uploading session",
])

# 4. Hybrid retrieval, with ablation results
add_slide("Hybrid Retrieval — Ablation Results", [
    "Dense (bge-m3) + sparse (lexical weights), fused with Reciprocal Rank Fusion, reranked by bge-reranker-v2-m3",
    "Measured on 20 real answerable questions (eval/questions.jsonl, facts copied from crawled text):",
    "  Dense only:        Recall@4 = 0.95,  MRR = 0.85",
    "  Hybrid (RRF):      Recall@4 = 0.95,  MRR = 0.642",
    "  Hybrid + rerank:   Recall@4 = 0.95,  MRR = 0.825",
    "RRF alone pushes the best passage down; reranking recovers it — and its 0-1 score is the evidence gate",
    "RRF scores are rank-based (rank 1 ≈ 1/61), so the old RRF threshold rejected almost nothing; the reranker gate (0.02) separates off-topic (~0.001) from answerable (≥ 0.25)",
])

# 5. Grounding and citations
add_slide("Grounding, Citations and Guardrails", [
    "Every fact cites a numbered source; sources are tagged official (te.eg) or uploaded, and are data, never instructions",
    "Validator drops unknown [S#] labels and flags numbers absent from the cited sources (list numbering and Arabic-Indic digits handled)",
    "Guardrails: small-talk shortcut → regex injection pre-filter → LLM router (off_topic / injection / question + clean standalone query) → reranker evidence gate",
    "Answer prompt gets only the standalone query (passing chat history made the 4B model answer one turn late)",
    "Output checks: wrong language/script → one regeneration; system-prompt leak → withheld",
    "Heavy test: 10/10 off-topic and 7/7 injection refused, 10/10 tricky legitimate questions answered; te.eg preferred over a conflicting upload, with both stated",
])

# 6. Speech: ASR comparison and TTS
add_slide("Speech: ASR Comparison and TTS", [
    "ASR: faster-whisper large-v3-turbo (deployed) vs large-v3, beam 1/5, telecom hotwords on/off — eval/asr_eval.py",
    "Synthetic TTS clips (plumbing check): deployed setting lowest WER (13.1%) and fastest; EN and MSA exact; large-v3 not better",
    "Hotword-echo guard: Whisper sometimes returned its own hotword prompt on clipped audio — detected and re-transcribed",
    "TTS: Piper, English (en_US-lessac-medium) + Arabic (ar_JO-kareem-medium); speech_text.py strips citations, verbalizes numbers, respells brand names",
    "Real recorded clips (Egyptian dialect, noise, code-switching) still to be collected — eval/RECORDING_CHECKLIST.md",
    "Limitation: Piper's Arabic voice is more robotic than English — a better Arabic/Egyptian voice is on the roadmap",
])

# 7. On-prem and offline deployment
add_slide("On-Prem, Offline, and Why CPU-Sized Models", [
    "HF_HUB_OFFLINE=1 / TRANSFORMERS_OFFLINE=1 everywhere; models load from local disk only; every call goes to 127.0.0.1",
    "Qwen3-4B Q4 chosen to keep CPU latency bounded; the same models run unchanged on a GPU",
    "CPU laptop: ~66 s p50 per answer.  RTX 3080: 1.2 s p50 — LLM, ASR, embedder and reranker on the GPU",
    "GPU memory (measured): LLM 3.96 GB, ASR 2.5 GB, embedder 1.9 GB, reranker +0.9 GB → 9.2 GB for the app; 24 GB cards run eval alongside",
    "Cloud (OpenRouter) is opt-in only, with a visible 'data leaves this machine' badge — never the default",
])

# 8. Evaluation results
add_slide("Evaluation Results", [
    "25 real questions (10 EN, 10 AR mixing MSA/Egyptian, 5 unanswerable) — facts copied from crawled text",
    "CPU baseline: 19/20 answered with citations, 2/5 unanswerable refused, p50 66.5 s / p95 100.7 s",
    "GPU + guardrails: 20/20 answered with citations, 5/5 refused, p50 1.2 s / p95 1.9 s",
    "8 simultaneous users: 32/32 answered, p50 5.8 s (after sizing llama-server's KV pool per slot)",
    "Heavy system test 49/49, 101 unit + 6 live tests, GPU soak test without memory growth",
    "Local vs OpenRouter comparison: not run — the provided API key was rejected (401)",
])

# 9. Limitations
add_slide("Limitations", [
    "Containers never run (no Docker daemon on either machine) — native processes used; Docker backend image uses CPU torch",
    "No real recorded audio yet — ASR accuracy measured on synthetic TTS speech only",
    "CPU latency ~1-2 min per answer; the router's extra LLM call on CPU not re-measured",
    "10 GB GPUs are tight: re-indexing or eval needs the backend stopped",
    "A conflicting upload is only flagged when the official passage is also retrieved",
    "4B-model quirks (own unit conversions, language drift) need the validator and reminders; Piper's Arabic voice lags English",
])

# 10. Production roadmap
add_slide("Production Roadmap", [
    "24 GB GPU (RTX 3090) and a larger model (Qwen3-8B) evaluated on the same question set; vLLM for many concurrent users",
    "Chatterbox or a dedicated Egyptian Arabic TTS voice",
    "PaddleOCR or a VLM-based OCR path for harder scanned documents",
    "Postgres instead of SQLite; queue workers for ingestion at scale",
    "SSO / RBAC instead of session-id scoping",
    "Auto-restarting services, CUDA backend image, monitoring/observability",
    "Whisper fine-tuning on real telecom-support audio",
    "Human handoff path for needs_escalation conversations",
])

os.makedirs(os.path.dirname(OUT_PATH), exist_ok=True)
prs.save(OUT_PATH)
print(f"wrote {OUT_PATH} ({len(prs.slides)} slides)")
